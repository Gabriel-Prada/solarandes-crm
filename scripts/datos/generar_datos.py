"""Genera los datos ficticios de SolarAndes.

Produce:
- data/generados/*.csv: los datos "verdaderos" del negocio, listos para cargar al CRM.
- data/heredado/clientes_heredado.xlsx: el Excel sucio que los asesores llevaban antes del CRM.
- data/heredado/respuestas_migracion.csv: qué contacto real hay detrás de cada fila del Excel,
  para medir qué tan bien funciona la migración de la fase 2.

Con la misma semilla, el resultado es siempre el mismo.

Uso: .venv/Scripts/python scripts/datos/generar_datos.py
"""

import csv
import random
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from faker import Faker
from openpyxl import Workbook
from openpyxl.styles import Font
from reglas import aplicar_descuento, cotizar

SEMILLA = 2026
HOY = date(2026, 10, 21)  # fecha fija para que los datos no cambien según el día en que se corran
NUM_CONTACTOS = 1000
NUM_CONTACTOS_HEREDADO = 280

RAIZ = Path(__file__).resolve().parents[2]
DIR_GENERADOS = RAIZ / "data" / "generados"
DIR_HEREDADO = RAIZ / "data" / "heredado"

# Dominios reservados para ejemplos: ninguna automatización le escribirá a una persona real.
DOMINIOS_CORREO = ["example.com", "example.net", "example.org"]

USUARIOS = [
    {"id": "U01", "nombre": "Laura Méndez", "rol": "gerente", "ciudad": "Bogotá"},
    {"id": "U02", "nombre": "Andrés Rojas", "rol": "asesor", "ciudad": "Bogotá"},
    {"id": "U03", "nombre": "Valentina Ospina", "rol": "asesor", "ciudad": "Medellín"},
    {"id": "U04", "nombre": "Camilo Torres", "rol": "asesor", "ciudad": "Cali"},
    {"id": "U05", "nombre": "Jhon Fredy Muñoz", "rol": "tecnico", "ciudad": ""},
    {"id": "U06", "nombre": "Marcela Díaz", "rol": "contadora", "ciudad": ""},
]
ASESOR_POR_CIUDAD = {"Bogotá": "U02", "Medellín": "U03", "Cali": "U04"}
TECNICO_ID = "U05"

PRODUCTOS = [
    {"sku": "PAN-550", "nombre": "Panel monocristalino 550 W", "categoria": "panel", "precio_cop": 620_000},
    {"sku": "INV-03", "nombre": "Inversor de red 3 kW", "categoria": "inversor", "precio_cop": 2_900_000},
    {"sku": "INV-05", "nombre": "Inversor de red 5 kW", "categoria": "inversor", "precio_cop": 4_100_000},
    {"sku": "INV-10", "nombre": "Inversor de red 10 kW", "categoria": "inversor", "precio_cop": 6_800_000},
    {"sku": "INV-20", "nombre": "Inversor trifásico 20 kW", "categoria": "inversor", "precio_cop": 11_500_000},
    {"sku": "BAT-05", "nombre": "Batería de litio 5 kWh", "categoria": "bateria", "precio_cop": 9_200_000},
    {"sku": "EST-01", "nombre": "Estructura de aluminio por panel", "categoria": "estructura", "precio_cop": 180_000},
    {"sku": "SRV-MNT", "nombre": "Mantenimiento preventivo", "categoria": "servicio", "precio_cop": 350_000},
]

PESOS_CIUDAD = {"Bogotá": 0.45, "Medellín": 0.30, "Cali": 0.25}
PESOS_CANAL = {"anuncios": 0.40, "whatsapp": 0.25, "web": 0.20, "referido": 0.15}
# Los referidos y los clientes que amplían su sistema cierran más; los de anuncios, menos.
MULTIPLICADOR_CANAL = {"referido": 1.15, "web": 1.0, "whatsapp": 1.0, "anuncios": 0.9, "ampliacion": 1.2}

PROBABILIDAD_ETAPA = {
    "nuevo": 5,
    "contactado": 15,
    "visita_agendada": 35,
    "cotizado": 50,
    "negociacion": 70,
    "ganado": 100,
    "perdido": 0,
}
# (desde, hacia, probabilidad base de avanzar, días mínimos y máximos que tarda)
TRANSICIONES = [
    ("nuevo", "contactado", 0.85, (0, 3)),
    ("contactado", "visita_agendada", 0.55, (2, 10)),
    ("visita_agendada", "cotizado", 0.80, (1, 4)),  # días contados desde la visita
    ("cotizado", "negociacion", 0.55, (3, 15)),
    ("negociacion", "ganado", 0.65, (3, 20)),
]
MOTIVOS_PERDIDA = {
    "temprano": {"no_respondio": 0.5, "aplazo": 0.2, "precio": 0.15, "sin_financiacion": 0.15},
    "visita_agendada": {"techo_no_apto": 0.35, "no_respondio": 0.25, "aplazo": 0.2, "precio": 0.2},
    "tarde": {"precio": 0.4, "competencia": 0.3, "sin_financiacion": 0.15, "aplazo": 0.15},
}
TEXTO_MOTIVO = {
    "precio": "le pareció caro",
    "competencia": "eligió a la competencia",
    "no_respondio": "dejó de responder",
    "techo_no_apto": "el techo no es apto",
    "sin_financiacion": "no consiguió financiación",
    "aplazo": "aplazó la decisión",
}

TIPOS_NEGOCIO = [
    "Panadería",
    "Restaurante",
    "Taller",
    "Ferretería",
    "Droguería",
    "Oficinas",
    "Lavandería",
    "Hotel",
    "Colegio",
    "Supermercado",
    "Carnicería",
    "Gimnasio",
]
CONSUMO_POR_ESTRATO = {2: (120, 200), 3: (150, 260), 4: (200, 350), 5: (280, 500), 6: (380, 750)}
PESOS_ESTRATO = {2: 0.05, 3: 0.30, 4: 0.35, 5: 0.20, 6: 0.10}
PREFIJOS_CELULAR = [
    "300",
    "301",
    "302",
    "304",
    "305",
    "310",
    "311",
    "312",
    "313",
    "314",
    "315",
    "316",
    "317",
    "318",
    "320",
    "321",
    "323",
    "350",
    "351",
]


def _elegir(rnd: random.Random, pesos: dict):
    return rnd.choices(list(pesos), weights=list(pesos.values()))[0]


def _sin_tildes(texto: str) -> str:
    descompuesto = unicodedata.normalize("NFD", texto)
    return "".join(c for c in descompuesto if unicodedata.category(c) != "Mn")


def _slug(texto: str) -> str:
    limpio = "".join(c if c.isalnum() else " " for c in _sin_tildes(texto).lower())
    return ".".join(limpio.split())


def _en_horario(rnd: random.Random, dia: date) -> datetime:
    return datetime.combine(dia, time(rnd.randint(8, 18), rnd.randint(0, 59)))


def _fmt(valor) -> str:
    if isinstance(valor, datetime):
        return valor.isoformat(sep=" ", timespec="minutes")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return "true" if valor else "false"
    return "" if valor is None else str(valor)


@dataclass
class Datos:
    usuarios: list[dict] = field(default_factory=list)
    productos: list[dict] = field(default_factory=list)
    contactos: list[dict] = field(default_factory=list)
    oportunidades: list[dict] = field(default_factory=list)
    actividades: list[dict] = field(default_factory=list)
    cotizaciones: list[dict] = field(default_factory=list)
    historial_etapas: list[dict] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Contactos
# ---------------------------------------------------------------------------


def _generar_contactos(rnd: random.Random, fake: Faker) -> list[dict]:
    contactos = []
    telefonos_usados, correos_usados = set(), set()

    for _ in range(NUM_CONTACTOS):
        tipo = "hogar" if rnd.random() < 0.75 else "negocio"
        ciudad = _elegir(rnd, PESOS_CIUDAD)
        nombre = f"{fake.first_name()} {fake.last_name()}"
        if rnd.random() < 0.5:
            nombre += f" {fake.last_name()}"

        if tipo == "hogar":
            empresa = ""
            estrato = _elegir(rnd, PESOS_ESTRATO)
            consumo = rnd.randint(*CONSUMO_POR_ESTRATO[estrato])
            base_correo = _slug(nombre.split()[0] + " " + nombre.split()[1])
        else:
            empresa = f"{rnd.choice(TIPOS_NEGOCIO)} {fake.last_name()}"
            if rnd.random() < 0.3:
                empresa += " S.A.S."
            estrato = None
            consumo = rnd.choice([rnd.randint(900, 1800), rnd.randint(1800, 3000), rnd.randint(3000, 4500)])
            base_correo = _slug(empresa.replace("S.A.S.", ""))

        telefono = None
        while telefono is None or telefono in telefonos_usados:
            telefono = "+57" + rnd.choice(PREFIJOS_CELULAR) + f"{rnd.randint(0, 9_999_999):07d}"
        telefonos_usados.add(telefono)

        correo = f"{base_correo}@{rnd.choice(DOMINIOS_CORREO)}"
        while correo in correos_usados:
            correo = f"{base_correo}{rnd.randint(1, 99)}@{rnd.choice(DOMINIOS_CORREO)}"
        correos_usados.add(correo)

        contactos.append(
            {
                "tipo": tipo,
                "nombre": nombre,
                "empresa": empresa,
                "telefono": telefono,
                "email": correo,
                "ciudad": ciudad,
                "direccion": fake.street_address(),
                "estrato": estrato,
                "consumo_kwh_mes": consumo,
                "canal_origen": _elegir(rnd, PESOS_CANAL),
                "asesor_id": ASESOR_POR_CIUDAD[ciudad],
                "creado_en": _en_horario(rnd, HOY - timedelta(days=rnd.randint(1, 365))),
            }
        )

    contactos.sort(key=lambda c: c["creado_en"])
    for i, c in enumerate(contactos, start=1):
        c["id"] = f"C{i:04d}"
    return contactos


# ---------------------------------------------------------------------------
# Oportunidades, actividades y cotizaciones
# ---------------------------------------------------------------------------


class _Simulador:
    """Recorre el embudo de ventas de una oportunidad y registra lo que pasa en cada paso."""

    def __init__(self, rnd: random.Random):
        self.rnd = rnd
        self.oportunidades: list[dict] = []
        self.actividades: list[dict] = []
        self.cotizaciones: list[dict] = []
        self.historial: list[dict] = []

    def _cambio_etapa(self, opp, anterior, nueva, fecha, cambiado_por):
        """Anota un paso del recorrido. cambiado_por=None significa que lo hizo el sistema."""
        self.historial.append(
            {
                "_opp": opp,
                "etapa_anterior": anterior,
                "etapa_nueva": nueva,
                "cambiado_por": cambiado_por,
                "fecha": fecha,
            }
        )

    def _actividad(self, opp, tipo, fecha, descripcion, usuario_id=None, completada=True):
        self.actividades.append(
            {
                "_opp": opp,
                "contacto_id": opp["contacto_id"],
                "usuario_id": usuario_id or opp["asesor_id"],
                "tipo": tipo,
                "descripcion": descripcion,
                "fecha": fecha,
                "completada": completada,
            }
        )

    def _cotizacion(self, opp, fecha, consumo, ciudad, descuento_pct, version):
        c = cotizar(consumo, ciudad)
        valor = aplicar_descuento(c.valor_total, descuento_pct)
        cot = {
            "_opp": opp,
            "version": version,
            "fecha": fecha,
            "consumo_kwh_mes": consumo,
            "potencia_kwp": c.potencia_instalada_kwp,
            "num_paneles": c.num_paneles,
            "precio_kwp": c.precio_kwp,
            "descuento_pct": descuento_pct,
            "valor_total": valor,
            "estado": "enviada",
        }
        self.cotizaciones.append(cot)
        return cot

    def simular(self, contacto: dict, creada: datetime, tipo: str, consumo: int) -> dict:
        rnd = self.rnd
        canal = "ampliacion" if tipo == "ampliacion" else contacto["canal_origen"]
        ciudad = contacto["ciudad"]
        estimado = cotizar(consumo, ciudad)
        opp = {
            "contacto_id": contacto["id"],
            "asesor_id": contacto["asesor_id"],
            "tipo": tipo,
            "nombre": f"{'Ampliación' if tipo == 'ampliacion' else 'Sistema'} "
            f"{estimado.potencia_instalada_kwp:g} kWp - {contacto['empresa'] or contacto['nombre']}",
            "etapa": "nuevo",
            "potencia_kwp": estimado.potencia_instalada_kwp,
            "valor_cop": estimado.valor_total,
            "probabilidad": 0,
            "motivo_perdida": "",
            "creada_en": creada,
            "actualizada_en": creada,
            "fecha_cierre": None,
        }
        self.oportunidades.append(opp)
        self._cambio_etapa(opp, None, "nuevo", creada, None)  # el lead entra solo, por su canal

        origen = "una solicitud de ampliación" if tipo == "ampliacion" else canal
        self._actividad(opp, "nota", creada, f"Lead recibido por {origen}.")

        t = creada.date()
        cotizaciones_opp = []
        fecha_visita = None
        estancada = False

        for desde, hacia, prob, (dmin, dmax) in TRANSICIONES:
            prob = min(0.95, prob * MULTIPLICADOR_CANAL[canal])
            if rnd.random() < prob:
                base = fecha_visita if desde == "visita_agendada" else t
                fecha = base + timedelta(days=rnd.randint(dmin, dmax))
                if fecha > HOY:
                    break  # el siguiente paso todavía no ha ocurrido
                t = fecha
                momento = max(_en_horario(rnd, t), opp["actualizada_en"] + timedelta(minutes=rnd.randint(15, 180)))
                self._cambio_etapa(opp, desde, hacia, momento, opp["asesor_id"])
                opp["etapa"], opp["actualizada_en"] = hacia, momento

                if hacia == "contactado":
                    medio = rnd.choice(["llamada", "whatsapp"])
                    self._actividad(
                        opp, medio, momento, f"Primer contacto. Consume {consumo} kWh/mes y quiere bajar la factura."
                    )
                elif hacia == "visita_agendada":
                    fecha_visita = t + timedelta(days=rnd.randint(2, 6))
                    agenda = f"Se agenda visita técnica para el {fecha_visita:%d/%m/%Y}."
                    self._actividad(opp, "llamada", momento, agenda)
                    self._actividad(
                        opp,
                        "visita",
                        _en_horario(rnd, fecha_visita),
                        "Visita técnica: revisar techo, sombras y tablero eléctrico.",
                        usuario_id=TECNICO_ID,
                        completada=fecha_visita <= HOY,
                    )
                elif hacia == "cotizado":
                    cot = self._cotizacion(opp, momento, consumo, ciudad, 0, 1)
                    cotizaciones_opp.append(cot)
                    self._actividad(
                        opp,
                        "correo",
                        momento,
                        f"Envío de cotización: {cot['num_paneles']} paneles, {cot['potencia_kwp']:g} kWp.",
                    )
                elif hacia == "negociacion":
                    pedido = rnd.choice(["un descuento", "financiación a 24 meses", "ajustar el número de paneles"])
                    self._actividad(opp, "llamada", momento, f"El cliente pide {pedido}.")
                    if rnd.random() < 0.55:
                        cot = self._cotizacion(opp, momento, consumo, ciudad, rnd.choice([3, 5, 5, 8]), 2)
                        cotizaciones_opp[-1]["estado"] = "reemplazada"
                        cotizaciones_opp.append(cot)
                        self._actividad(
                            opp,
                            "correo",
                            momento,
                            f"Envío de cotización ajustada con {cot['descuento_pct']} % de descuento.",
                        )
                elif hacia == "ganado":
                    opp["fecha_cierre"] = momento
                    self._actividad(opp, "nota", momento, "Firmó contrato y pagó el anticipo del 50 %.")
                    fecha_inst = t + timedelta(days=rnd.randint(7, 20))
                    self._actividad(
                        opp,
                        "tarea",
                        _en_horario(rnd, fecha_inst),
                        "Instalación del sistema.",
                        usuario_id=TECNICO_ID,
                        completada=fecha_inst <= HOY,
                    )
            else:
                if rnd.random() < 0.06:
                    estancada = True  # nadie la cerró: queda abierta y sin seguimiento
                    break
                minimo = (fecha_visita - t).days + 1 if desde == "visita_agendada" else 5
                fecha = t + timedelta(days=rnd.randint(minimo, max(minimo, 30)))
                if fecha > HOY:
                    break  # todavía no se sabe que se perdió
                grupo = (
                    desde
                    if desde == "visita_agendada"
                    else ("temprano" if desde in ("nuevo", "contactado") else "tarde")
                )
                motivo = _elegir(rnd, MOTIVOS_PERDIDA[grupo])
                momento = _en_horario(rnd, fecha)
                self._cambio_etapa(opp, desde, "perdido", momento, opp["asesor_id"])
                opp.update(etapa="perdido", motivo_perdida=motivo, actualizada_en=momento, fecha_cierre=momento)
                self._actividad(opp, "nota", momento, f"Oportunidad perdida: {TEXTO_MOTIVO[motivo]}.")
            if opp["etapa"] in ("ganado", "perdido"):
                break

        abierta = opp["etapa"] not in ("ganado", "perdido")
        if abierta and not estancada:
            dias_sin_actividad = (HOY - t).days
            for _ in range(rnd.randint(0, 2)):
                if dias_sin_actividad > 2:
                    dia = t + timedelta(days=rnd.randint(1, dias_sin_actividad - 1))
                    self._actividad(
                        opp,
                        rnd.choice(["llamada", "whatsapp"]),
                        _en_horario(rnd, dia),
                        "Seguimiento: el cliente sigue evaluando.",
                    )
            if rnd.random() < 0.6:
                dia = HOY + timedelta(days=rnd.randint(1, 5))
                self._actividad(opp, "tarea", _en_horario(rnd, dia), "Hacer seguimiento.", completada=False)

        if cotizaciones_opp:
            ultima = cotizaciones_opp[-1]
            ultima["estado"] = {"ganado": "aceptada", "perdido": "rechazada"}.get(opp["etapa"], "enviada")
            opp["valor_cop"] = ultima["valor_total"]
            opp["potencia_kwp"] = ultima["potencia_kwp"]
        opp["probabilidad"] = PROBABILIDAD_ETAPA[opp["etapa"]]
        return opp


def _generar_ventas(rnd: random.Random, contactos: list[dict], datos: Datos) -> None:
    sim = _Simulador(rnd)
    for contacto in contactos:
        opp = sim.simular(contacto, contacto["creado_en"], "nueva", contacto["consumo_kwh_mes"])
        # Algunos clientes satisfechos amplían su sistema meses después
        cierre = opp["fecha_cierre"]
        if opp["etapa"] == "ganado" and cierre.date() <= HOY - timedelta(days=150) and rnd.random() < 0.10:
            inicio = cierre.date() + timedelta(days=rnd.randint(120, 300))
            if inicio < HOY:
                extra = round(contacto["consumo_kwh_mes"] * rnd.uniform(0.3, 0.6))
                sim.simular(contacto, _en_horario(rnd, inicio), "ampliacion", extra)

    # Los identificadores se asignan en orden cronológico, como los asignaría un sistema real
    sim.oportunidades.sort(key=lambda o: (o["creada_en"], o["contacto_id"]))
    for i, o in enumerate(sim.oportunidades, start=1):
        o["id"] = f"O{i:04d}"

    sim.cotizaciones.sort(key=lambda c: (c["fecha"], c["_opp"]["id"], c["version"]))
    consecutivo_por_anio: dict[int, int] = {}
    for i, c in enumerate(sim.cotizaciones, start=1):
        anio = c["fecha"].year
        consecutivo_por_anio[anio] = consecutivo_por_anio.get(anio, 0) + 1
        c["id"] = f"Q{i:04d}"
        c["numero"] = f"COT-{anio}-{consecutivo_por_anio[anio]:04d}"
        c["oportunidad_id"] = c.pop("_opp")["id"]

    sim.actividades.sort(key=lambda a: (a["fecha"], a["_opp"]["id"]))
    for i, a in enumerate(sim.actividades, start=1):
        a["id"] = f"A{i:05d}"
        a["oportunidad_id"] = a.pop("_opp")["id"]

    sim.historial.sort(key=lambda h: (h["fecha"], h["_opp"]["id"]))
    for i, h in enumerate(sim.historial, start=1):
        h["id"] = f"H{i:05d}"
        h["oportunidad_id"] = h.pop("_opp")["id"]

    datos.oportunidades = sim.oportunidades
    datos.cotizaciones = sim.cotizaciones
    datos.actividades = sim.actividades
    datos.historial_etapas = sim.historial


def generar(semilla: int = SEMILLA) -> Datos:
    rnd = random.Random(semilla)
    fake = Faker("es_CO")
    fake.seed_instance(semilla)

    datos = Datos()
    datos.usuarios = [
        {**u, "email": f"{_slug(u['nombre'].split()[0] + ' ' + u['nombre'].split()[-1])}@solarandes.test"}
        for u in USUARIOS
    ]
    datos.productos = [dict(p) for p in PRODUCTOS]
    datos.contactos = _generar_contactos(rnd, fake)
    _generar_ventas(rnd, datos.contactos, datos)
    return datos


# ---------------------------------------------------------------------------
# Excel heredado (sucio a propósito)
# ---------------------------------------------------------------------------

HOJA_GENERAL = "Clientes BOG-CALI"
HOJA_MEDELLIN = "Medellín (Vale)"
COLUMNAS_GENERAL = [
    "Nombre completo",
    "Cel",
    "Correo",
    "Ciudad",
    "Tipo cliente",
    "Consumo (kWh)",
    "Asesor",
    "Fecha contacto",
    "Estado",
    "Notas",
]
# La hoja de Valentina tiene otro orden, otros nombres y no tiene ciudad ni asesor
COLUMNAS_MEDELLIN = ["FECHA", "CLIENTE", "TELEFONO", "EMAIL", "CONSUMO MES", "TIPO", "ESTADO", "OBSERVACIONES"]
CAMPO_POR_COLUMNA = {
    "Nombre completo": "nombre",
    "Cel": "telefono",
    "Correo": "correo",
    "Ciudad": "ciudad",
    "Tipo cliente": "tipo",
    "Consumo (kWh)": "consumo",
    "Asesor": "asesor",
    "Fecha contacto": "fecha",
    "Estado": "estado",
    "Notas": "notas",
    "FECHA": "fecha",
    "CLIENTE": "nombre",
    "TELEFONO": "telefono",
    "EMAIL": "correo",
    "CONSUMO MES": "consumo",
    "TIPO": "tipo",
    "ESTADO": "estado",
    "OBSERVACIONES": "notas",
}

VARIANTES_CIUDAD = {
    "Bogotá": ["Bogotá", "Bogota", "BOGOTA", "bogotá", "Bogotá D.C.", "Bogota DC", "Bta", "Bogotá "],
    "Medellín": ["Medellín", "Medellin", "MEDELLIN", "medellin", "Medellín, Antioquia", "Mde"],
    "Cali": ["Cali", "CALI", "cali", "Cali - Valle", "Santiago de Cali", "Cali "],
}
VARIANTES_TIPO = {
    "hogar": ["Hogar", "hogar", "Residencial", "Casa", "RESIDENCIAL"],
    "negocio": ["Negocio", "Comercial", "Empresa", "negocio", "COMERCIAL"],
}
VARIANTES_ASESOR = {
    "U02": ["Andrés", "andres rojas", "A. Rojas", "Andres R", "ANDRÉS"],
    "U03": ["Vale", "Valentina O.", "valentina"],
    "U04": ["Camilo", "camilo t", "C. Torres", "Cami"],
}
VARIANTES_ESTADO = {
    "nuevo": ["nuevo", "Nuevo lead", "sin contactar"],
    "contactado": ["contactado", "llamado", "Interesado", "interesado!!"],
    "visita_agendada": ["visita pendiente", "Visita agendada", "VISITA"],
    "cotizado": ["cotizado", "Cotizado ✓", "se envió cotización", "COT enviada"],
    "negociacion": ["negociando", "pide descuento", "en negociación"],
    "ganado": ["cerró!!", "GANADO", "vendido", "Instalado"],
    "perdido": ["PERDIDO", "no le interesó", "no contesta"],
}
MESES = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
]


def _ensuciar_nombre(rnd, contacto) -> str:
    nombre = contacto["nombre"]
    if contacto["empresa"]:
        nombre = rnd.choice([contacto["empresa"], f"{contacto['empresa']} - {nombre}", nombre])
    elif rnd.random() < 0.10:
        partes = nombre.split()
        nombre = " ".join(partes[1:] + partes[:1])  # "Apellido Nombre"
    if rnd.random() < 0.30:
        nombre = _sin_tildes(nombre)
    r = rnd.random()
    if r < 0.25:
        nombre = nombre.upper()
    elif r < 0.40:
        nombre = nombre.lower()
    if rnd.random() < 0.10:
        nombre = nombre.replace(" ", "  ", 1)
    if rnd.random() < 0.20:
        nombre += " "
    return nombre


def _ensuciar_telefono(rnd, telefono) -> tuple[object, bool]:
    """Devuelve el teléfono sucio y si todavía se puede recuperar."""
    d = telefono[3:]  # los 10 dígitos sin +57
    r = rnd.random()
    if r < 0.05:
        return None, False
    if r < 0.08:
        return d[:4] + d[5:], False  # le falta un dígito
    formatos = [
        d,
        d,
        int(d),
        f"{d}.0",
        f"{d[:3]} {d[3:6]} {d[6:]}",
        f"{d[:3]}-{d[3:6]}-{d[6:]}",
        f"+57 {d[:3]} {d[3:]}",
        f"57{d}",
        f"(+57) {d[:3]}-{d[3:6]}-{d[6:8]}-{d[8:]}",
    ]
    return rnd.choice(formatos), True


def _ensuciar_correo(rnd, correo) -> tuple[object, bool]:
    r = rnd.random()
    if r < 0.10:
        return None, False
    if r < 0.15:
        return correo.rsplit(".", 1)[0], False  # sin ".com"
    if r < 0.30:
        return correo.upper(), True
    if r < 0.40:
        return f" {correo} ", True
    return correo, True


def _ensuciar_consumo(rnd, consumo):
    return rnd.choices(
        [consumo, f"{consumo} kWh", f"{consumo}kwh", f"aprox {consumo}", float(consumo), None],
        weights=[35, 20, 10, 10, 15, 10],
    )[0]


def _ensuciar_fecha(rnd, momento: datetime):
    d = momento.date()
    return rnd.choices(
        [
            d.strftime("%d/%m/%Y"),
            d.isoformat(),
            datetime.combine(d, time()),
            f"{d.day} de {MESES[d.month - 1]} de {d.year}",
            f"{d.day}-{MESES[d.month - 1][:3]}-{d:%y}",
            None,
        ],
        weights=[30, 15, 30, 10, 10, 5],
    )[0]


def _nota(rnd, contacto, contactos) -> str | None:
    if rnd.random() > 0.25:
        return None
    referido = rnd.choice(contactos)["nombre"].split()[0]
    return rnd.choice(
        [
            "Llamar después de las 5 pm",
            f"Tel fijo 60{rnd.choice('1426')} {rnd.randint(2000000, 7999999)}",
            f"Referido por {referido}",
            "Pidió financiación",
            "Escribir por WhatsApp, no contesta llamadas",
            "Tiene techo de teja, revisar",
            f"Correo alterno: {_slug(contacto['nombre'].split()[0])}@example.net",
        ]
    )


def construir_heredado(rnd: random.Random, datos: Datos) -> dict[str, list[tuple[list, dict]]]:
    """Arma las hojas del Excel. Cada fila va con sus respuestas correctas."""
    primera_opp = {}
    for o in datos.oportunidades:
        primera_opp.setdefault(o["contacto_id"], o)

    elegidos = datos.contactos[:NUM_CONTACTOS_HEREDADO]  # los más antiguos
    filas = {HOJA_GENERAL: [], HOJA_MEDELLIN: []}

    for contacto in elegidos:
        copias = rnd.choices([1, 2, 3], weights=[82, 14, 4])[0]
        for _ in range(copias):
            medellin = contacto["ciudad"] == "Medellín"
            hoja = HOJA_MEDELLIN if medellin and rnd.random() < 0.85 else HOJA_GENERAL
            telefono, tel_ok = _ensuciar_telefono(rnd, contacto["telefono"])
            correo, correo_ok = _ensuciar_correo(rnd, contacto["email"])
            campos = {
                "nombre": _ensuciar_nombre(rnd, contacto),
                "telefono": telefono,
                "correo": correo,
                "ciudad": None if rnd.random() < 0.05 else rnd.choice(VARIANTES_CIUDAD[contacto["ciudad"]]),
                "tipo": None if rnd.random() < 0.05 else rnd.choice(VARIANTES_TIPO[contacto["tipo"]]),
                "consumo": _ensuciar_consumo(rnd, contacto["consumo_kwh_mes"]),
                "asesor": None if rnd.random() < 0.05 else rnd.choice(VARIANTES_ASESOR[contacto["asesor_id"]]),
                "fecha": _ensuciar_fecha(rnd, contacto["creado_en"]),
                "estado": rnd.choice(VARIANTES_ESTADO[primera_opp[contacto["id"]]["etapa"]]),
                "notas": _nota(rnd, contacto, datos.contactos),
            }
            columnas = COLUMNAS_MEDELLIN if hoja == HOJA_MEDELLIN else COLUMNAS_GENERAL
            valores = [campos[CAMPO_POR_COLUMNA[c]] for c in columnas]
            respuesta = {
                "tipo_fila": "dato",
                "contacto_id": contacto["id"],
                "telefono_recuperable": tel_ok,
                "correo_recuperable": correo_ok,
            }
            filas[hoja].append((valores, respuesta))

    def basura(tipo_fila, valores=None):
        return (
            valores or [],
            {"tipo_fila": tipo_fila, "contacto_id": "", "telefono_recuperable": "", "correo_recuperable": ""},
        )

    hojas = {}
    general = filas[HOJA_GENERAL]
    rnd.shuffle(general)
    general.insert(len(general) // 2, basura("encabezado_repetido", list(COLUMNAS_GENERAL)))
    for _ in range(3):
        general.insert(rnd.randint(1, len(general) - 1), basura("vacia"))
    total_datos = sum(1 for _, r in general if r["tipo_fila"] == "dato")
    general.append(basura("total", ["TOTAL CLIENTES", total_datos]))
    hojas[HOJA_GENERAL] = [basura("encabezado", list(COLUMNAS_GENERAL)), *general]

    medellin = filas[HOJA_MEDELLIN]
    rnd.shuffle(medellin)
    for _ in range(2):
        medellin.insert(rnd.randint(1, len(medellin) - 1), basura("vacia"))
    hojas[HOJA_MEDELLIN] = [
        basura("titulo", ["Clientes Medellín - Valentina (no borrar)"]),
        basura("encabezado", list(COLUMNAS_MEDELLIN)),
        *medellin,
    ]
    return hojas


# ---------------------------------------------------------------------------
# Escritura de archivos
# ---------------------------------------------------------------------------

COLUMNAS_CSV = {
    "usuarios": ["id", "nombre", "email", "rol", "ciudad"],
    "productos": ["sku", "nombre", "categoria", "precio_cop"],
    "contactos": [
        "id",
        "tipo",
        "nombre",
        "empresa",
        "telefono",
        "email",
        "ciudad",
        "direccion",
        "estrato",
        "consumo_kwh_mes",
        "canal_origen",
        "asesor_id",
        "creado_en",
    ],
    "oportunidades": [
        "id",
        "contacto_id",
        "asesor_id",
        "tipo",
        "nombre",
        "etapa",
        "potencia_kwp",
        "valor_cop",
        "probabilidad",
        "motivo_perdida",
        "creada_en",
        "actualizada_en",
        "fecha_cierre",
    ],
    "actividades": ["id", "oportunidad_id", "contacto_id", "usuario_id", "tipo", "descripcion", "fecha", "completada"],
    "cotizaciones": [
        "id",
        "numero",
        "oportunidad_id",
        "version",
        "fecha",
        "consumo_kwh_mes",
        "potencia_kwp",
        "num_paneles",
        "precio_kwp",
        "descuento_pct",
        "valor_total",
        "estado",
    ],
    "historial_etapas": ["id", "oportunidad_id", "etapa_anterior", "etapa_nueva", "cambiado_por", "fecha"],
}


def escribir_csv(ruta: Path, filas: list[dict], columnas: list[str]) -> None:
    with ruta.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.writer(f)
        escritor.writerow(columnas)
        for fila in filas:
            escritor.writerow([_fmt(fila[c]) for c in columnas])


def escribir_heredado(hojas: dict, ruta_excel: Path, ruta_respuestas: Path) -> None:
    libro = Workbook()
    libro.remove(libro.active)
    respuestas = []
    for nombre_hoja, filas in hojas.items():
        hoja = libro.create_sheet(nombre_hoja)
        for num_fila, (valores, respuesta) in enumerate(filas, start=1):
            for num_col, valor in enumerate(valores, start=1):
                celda = hoja.cell(row=num_fila, column=num_col, value=valor)
                if isinstance(valor, datetime):
                    celda.number_format = "dd/mm/yyyy"
                if respuesta["tipo_fila"] in ("encabezado", "titulo"):
                    celda.font = Font(bold=True)
            respuestas.append({"hoja": nombre_hoja, "fila": num_fila, **respuesta})
        for letra in "ABCDEFGHIJ":
            hoja.column_dimensions[letra].width = 22
    libro.save(ruta_excel)
    escribir_csv(
        ruta_respuestas,
        respuestas,
        ["hoja", "fila", "tipo_fila", "contacto_id", "telefono_recuperable", "correo_recuperable"],
    )


def main(dir_generados: Path = DIR_GENERADOS, dir_heredado: Path = DIR_HEREDADO) -> Datos:
    datos = generar()
    dir_generados.mkdir(parents=True, exist_ok=True)
    dir_heredado.mkdir(parents=True, exist_ok=True)
    for nombre, columnas in COLUMNAS_CSV.items():
        escribir_csv(dir_generados / f"{nombre}.csv", getattr(datos, nombre), columnas)

    hojas = construir_heredado(random.Random(SEMILLA + 1), datos)
    escribir_heredado(hojas, dir_heredado / "clientes_heredado.xlsx", dir_heredado / "respuestas_migracion.csv")

    etapas: dict[str, int] = {}
    for o in datos.oportunidades:
        etapas[o["etapa"]] = etapas.get(o["etapa"], 0) + 1
    print(
        f"Contactos: {len(datos.contactos)} | Oportunidades: {len(datos.oportunidades)} | "
        f"Actividades: {len(datos.actividades)} | Cotizaciones: {len(datos.cotizaciones)}"
    )
    print("Por etapa:", ", ".join(f"{e}={n}" for e, n in sorted(etapas.items())))
    print(f"Excel heredado: {sum(len(f) for f in hojas.values())} filas en {len(hojas)} hojas")
    return datos


if __name__ == "__main__":
    main()
