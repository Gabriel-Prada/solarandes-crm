"""Paso 1 de la migración: leer el Excel heredado y dejarlo como una sola tabla.

El Excel tiene dos hojas con columnas distintas, un título encima del encabezado en una de
ellas, filas vacías, un encabezado repetido en la mitad y una fila de "TOTAL". Este módulo
no limpia los valores (eso es el paso 2): solo encuentra las filas que son personas y les
pone a todas los mismos nombres de columna.
"""

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
RUTA_EXCEL = RAIZ / "data" / "heredado" / "clientes_heredado.xlsx"

# Cada nombre de columna que se ha visto en el Excel, ya sin tildes ni mayúsculas, y la
# columna estándar a la que corresponde
COLUMNAS = {
    "nombre completo": "nombre",
    "cliente": "nombre",
    "cel": "telefono",
    "telefono": "telefono",
    "correo": "correo",
    "email": "correo",
    "ciudad": "ciudad",
    "tipo cliente": "tipo",
    "tipo": "tipo",
    "consumo (kwh)": "consumo",
    "consumo mes": "consumo",
    "asesor": "asesor",
    "fecha contacto": "fecha",
    "fecha": "fecha",
    "estado": "estado",
    "notas": "notas",
    "observaciones": "notas",
}


@dataclass
class FilaHeredada:
    """Una fila del Excel que corresponde a una persona, con los valores tal como venían."""

    hoja: str
    fila: int  # número de fila en el Excel (la primera es 1), para poder ir a buscarla
    valores: dict[str, object] = field(default_factory=dict)


def simplificar(texto: object) -> str:
    """Minúsculas, sin tildes y sin espacios de sobra: "  Teléfono " → "telefono"."""
    sin_tildes = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.lower().split())


def esta_vacia(celdas: tuple) -> bool:
    return all(c is None or str(c).strip() == "" for c in celdas)


def columnas_de_encabezado(celdas: tuple) -> dict[int, str] | None:
    """Si la fila es un encabezado, devuelve {posición: columna estándar}; si no, None.

    Se considera encabezado cuando al menos 3 celdas son nombres de columna conocidos: así
    un dato que casualmente diga "Cliente" no se confunde con un encabezado.
    """
    posiciones = {i: COLUMNAS[simplificar(c)] for i, c in enumerate(celdas) if c and simplificar(c) in COLUMNAS}
    return posiciones if len(posiciones) >= 3 else None


def es_fila_de_total(celdas: tuple) -> bool:
    """ "TOTAL CLIENTES" o "Total:" sí; "Totalgas S.A.S." no: "total" tiene que ser una palabra suelta."""
    primera = next((c for c in celdas if c is not None and str(c).strip()), "")
    return re.match(r"total\b", simplificar(primera)) is not None


def leer_excel(ruta: Path = RUTA_EXCEL) -> list[FilaHeredada]:
    libro = load_workbook(ruta, read_only=True)
    filas: list[FilaHeredada] = []
    for hoja in libro.worksheets:
        columnas: dict[int, str] | None = None
        for numero, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
            if esta_vacia(celdas):
                continue
            encabezado = columnas_de_encabezado(celdas)
            if encabezado:
                columnas = encabezado  # el primero, o uno repetido en la mitad: se salta igual
                continue
            if columnas is None or es_fila_de_total(celdas):
                continue  # un título antes del encabezado, o el total del final
            valores = {nombre: celdas[i] if i < len(celdas) else None for i, nombre in columnas.items()}
            filas.append(FilaHeredada(hoja=hoja.title, fila=numero, valores=valores))
    libro.close()
    return filas
