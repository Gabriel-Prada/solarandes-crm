"""Paso 5 de la migración: el reporte de calidad.

Genera dos archivos:
- docs/reporte-migracion.md: qué entró, qué se limpió, qué se unió, qué se cargó al CRM y
  cómo sabemos que está bien.
- data/migracion/para_revisar.csv: lo que tiene que decidir una persona, para abrir en Excel.

Uso: .venv/Scripts/python scripts/migracion/reporte.py
Lee el CRM (no escribe nada). Si llega otro Excel, se vuelve a correr y el reporte se actualiza.
"""

import csv
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path

sys.path[:0] = [str(Path(__file__).resolve().parents[1] / d) for d in ("supabase", "datos")]

from conciliar import PREFIJO_NOTA, ContactoCRM, EnRevision, Plan, planear  # noqa: E402
from duplicados import (  # noqa: E402
    ParejaEnRevision,
    Persona,
    agrupar,
    comparar_nombres,
    contradicciones,
    deduplicar,
)
from leer_excel import RAIZ, RUTA_EXCEL, leer_excel  # noqa: E402
from normalizar import ContactoLimpio, normalizar  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RUTA_REPORTE = RAIZ / "docs" / "reporte-migracion.md"
RUTA_REVISION = RAIZ / "data" / "migracion" / "para_revisar.csv"
RUTA_RESPUESTAS = RAIZ / "data" / "heredado" / "respuestas_migracion.csv"
RUTA_VERDAD = RAIZ / "data" / "generados" / "contactos.csv"

# Campo de la fila limpia → cómo se llama en el reporte
CAMPOS = {
    "telefono": "Teléfono", "correo": "Correo", "ciudad": "Ciudad", "tipo": "Tipo de cliente",
    "consumo_kwh_mes": "Consumo", "asesor": "Asesor", "fecha_contacto": "Fecha", "etapa": "Estado",
}  # fmt: skip


@dataclass
class DatosReporte:
    filas_en_excel: int  # todas, incluidas las que no son personas
    contactos: list[ContactoLimpio]  # las filas que son personas, ya limpias
    personas: list[Persona]
    revision_duplicados: list[ParejaEnRevision]
    plan: Plan
    notas_en_crm: int  # notas del Excel que ya están guardadas en el CRM
    crm: list[ContactoCRM]
    revision_crm: list[ParejaEnRevision]  # parejas del CRM con el mismo nombre (homónimos)
    grupos_duplicados_crm: int
    fecha: date


def contar_filas_excel(ruta: Path = RUTA_EXCEL) -> int:
    libro = load_workbook(ruta, read_only=True)
    total = sum(1 for hoja in libro.worksheets for _ in hoja.iter_rows())
    libro.close()
    return total


def limpieza_por_campo(contactos: list[ContactoLimpio]) -> list[tuple[str, int, int, int, int]]:
    """Por campo: (nombre, entendido, de ellos deducido de la hoja, escrito pero ilegible, sin dato)."""
    problemas = Counter(p for c in contactos for p in c.problemas)
    return [
        (nombre, sum(getattr(c, campo) is not None for c in contactos), problemas[f"{campo}_de_la_hoja"],
         problemas[f"{campo}_ilegible"], problemas[f"{campo}_vacio"])
        for campo, nombre in CAMPOS.items()
    ]  # fmt: skip


def nombre_de(persona: Persona) -> str:
    return persona.nombre or persona.empresa or "sin nombre"


def casos_del_plan(plan: Plan) -> list[EnRevision]:
    """Los casos en revisión del paso 4, sin los que ya vienen de una duda del paso 3.

    Una pareja en duda del paso 3 marca a sus dos personas como pendientes, y el paso 4 las
    manda a revisión otra vez: contarlas de nuevo sería contar tres veces la misma decisión.
    """
    return [c for c in plan.revision if "pendiente_de_revision" not in c.persona.problemas]


def para_revisar(datos: DatosReporte) -> list[dict[str, str]]:
    """Cada decisión que tiene que tomar una persona, con lo necesario para buscarla en el Excel."""
    filas = []
    for pareja in datos.revision_duplicados:
        filas.append({
            "tipo": "posible duplicado en el Excel",
            "filas del Excel": donde([(pareja.a.hoja, pareja.a.fila), (pareja.b.hoja, pareja.b.fila)]),
            "nombre": pareja.a.nombre or pareja.a.empresa or "sin nombre",
            "motivo": pareja.razon,
        })  # fmt: skip
    for caso in casos_del_plan(datos.plan):
        filas.append({
            "tipo": "no se pudo ubicar en el CRM",
            "filas del Excel": donde(caso.persona.filas),
            "nombre": nombre_de(caso.persona),
            "motivo": caso.razon,
        })  # fmt: skip
    return filas


def separar_dudas_del_crm(revision_crm: list[ParejaEnRevision]) -> tuple[list, list]:
    """(homónimos, posibles duplicados). Homónimo: mismo nombre y datos que se contradicen, o sea,
    otra persona que se llama igual. Lo demás es una duda de verdad y no se presenta como resuelta."""
    homonimos, posibles = [], []
    for pareja in revision_crm:
        mismo_nombre = comparar_nombres(pareja.a, pareja.b)[0]
        (homonimos if mismo_nombre and contradicciones(pareja.a, pareja.b) else posibles).append(pareja)
    return homonimos, posibles


def calificar(datos: DatosReporte, ruta_respuestas: Path = RUTA_RESPUESTAS, ruta_verdad: Path = RUTA_VERDAD) -> dict:
    """Compara contra el archivo de respuestas. Solo existe porque los datos son ficticios."""
    with ruta_respuestas.open(encoding="utf-8") as f:
        respuesta_de = {(r["hoja"], int(r["fila"])): r for r in csv.DictReader(f)}
    with ruta_verdad.open(encoding="utf-8") as f:
        verdad = {c["id"]: c for c in csv.DictReader(f)}
    columnas = {"telefono": "telefono", "correo": "email", "ciudad": "ciudad", "tipo": "tipo", "asesor": "asesor_id"}

    equivocados = 0
    recuperados = {"telefono": [0, 0], "correo": [0, 0]}  # [recuperados, recuperables], fila por fila
    for c in datos.contactos:
        respuesta = respuesta_de[(c.hoja, c.fila)]
        real = verdad[respuesta["contacto_id"]]
        equivocados += sum(getattr(c, campo) not in (None, real[col]) for campo, col in columnas.items())
        for campo, cuenta in recuperados.items():
            if respuesta[f"{campo}_recuperable"] == "true":
                cuenta[1] += 1
                cuenta[0] += getattr(c, campo) is not None

    reales_por_persona = [{respuesta_de[f]["contacto_id"] for f in p.filas} for p in datos.personas]
    fichas_por_real = Counter(real for reales in reales_por_persona for real in reales)
    return {
        "equivocados": equivocados,
        "telefonos": tuple(recuperados["telefono"]),
        "correos": tuple(recuperados["correo"]),
        "mezcladas": sum(len(reales) > 1 for reales in reales_por_persona),
        "repartidas": sum(veces > 1 for veces in fichas_por_real.values()),
        "bien_asociadas": sum(
            e.contacto.codigo == respuesta_de[e.persona.filas[0]]["contacto_id"] for e in datos.plan.existentes
        ),
    }


def markdown(datos: DatosReporte, nota: dict | None = None) -> str:
    filas_personas = len(datos.contactos)
    personas = len(datos.personas)
    con_repetidas = sum(len(p.filas) > 1 for p in datos.personas)
    plan = datos.plan
    casos = len(para_revisar(datos))
    faltantes = {campo: sum(getattr(p, campo) is None for p in datos.personas) for campo in CAMPOS}
    descartadas = datos.filas_en_excel - filas_personas
    notas_pendientes = sum(len(e.notas_nuevas) for e in plan.existentes)
    datos_pendientes = sum(len(e.completar) for e in plan.existentes)
    homonimos, posibles = separar_dudas_del_crm(datos.revision_crm)

    if datos.grupos_duplicados_crm:
        estado_crm = f"Ojo: el CRM tiene {datos.grupos_duplicados_crm} grupo(s) de duplicados."
    elif posibles:
        estado_crm = f"El CRM no tiene duplicados seguros, pero hay {len(posibles)} pareja(s) por revisar."
    else:
        estado_crm = "El CRM no tiene duplicados."
    pendiente = notas_pendientes or datos_pendientes or plan.nuevas
    aviso = (
        [
            "",
            f"**Ojo: la migración no está al día.** Faltan {notas_pendientes} nota(s), {datos_pendientes} dato(s) y "
            f"{len(plan.nuevas)} persona(s) nueva(s). Se guardan con `scripts/migracion/migrar.py --aplicar`.",
        ]
        if pendiente
        else []
    )

    md = [
        "# Reporte de la migración del Excel heredado",
        "",
        f"Generado el {datos.fecha:%Y-%m-%d} con `scripts/migracion/reporte.py`. Todos los datos son ficticios.",
        "",
        "## En una frase",
        "",
        f"El Excel tenía {filas_personas} filas de clientes, que resultaron ser {personas} personas. "
        f"{len(plan.existentes)} ya están en el CRM, con {datos.notas_en_crm} notas que antes solo vivían en el "
        f"Excel; {casos} caso(s) quedan para que los decida una persona. {estado_crm}",
        *aviso,
        "",
        "| | |",
        "|---|---|",
        f"| Filas en el Excel | {datos.filas_en_excel} |",
        f"| Filas que no eran clientes (título, encabezados, vacías, total) | {descartadas} |",
        f"| Filas de clientes | {filas_personas} |",
        f"| Personas distintas, después de unir duplicados | {personas} |",
        f"| Duplicados unidos | {filas_personas - personas} filas, en {con_repetidas} personas |",
        f"| Personas que ya están en el CRM | {len(plan.existentes)} |",
        f"| Personas nuevas pendientes de crear | {len(plan.nuevas)} |",
        f"| Notas del Excel guardadas en el CRM | {datos.notas_en_crm} |",
        f"| Notas pendientes de guardar | {notas_pendientes} |",
        f"| Casos para revisar | {casos} |",
        f"| Duplicados en el CRM | {datos.grupos_duplicados_crm} |",
        "",
        "## 1. Limpieza, campo por campo",
        "",
        "Regla: la limpieza nunca inventa datos. Si un valor no se entiende, queda vacío y se anota.",
        "",
        "| Campo | Entendido | De ellos, deducido de la hoja | Escrito pero ilegible | Sin dato |",
        "|---|---|---|---|---|",
    ]
    for nombre, entendido, de_la_hoja, ilegible, vacio in limpieza_por_campo(datos.contactos):
        md.append(f"| {nombre} | {entendido} de {filas_personas} | {de_la_hoja} | {ilegible} | {vacio} |")
    md += [
        "",
        '"Deducido de la hoja": la hoja "Medellín (Vale)" no tiene columnas de ciudad ni asesor. Si la celda '
        "está vacía, se asume lo que dice el nombre de la hoja (Medellín, Valentina). Es una suposición razonable, "
        "y por eso se cuenta aparte.",
        "",
        "## 2. Duplicados",
        "",
        f"{filas_personas} filas → {personas} personas. Dos filas se unen solas si comparten teléfono o correo, "
        "o el mismo nombre y la misma ciudad, y nada se contradice. Ante la duda, van a revisión: unir a dos "
        "personas distintas es peor que dejar un duplicado.",
        "",
        f"- Personas que aparecían más de una vez: {con_repetidas}",
        f"- Parejas que quedaron en duda: {len(datos.revision_duplicados)}",
        "",
        "Unir filas repetidas también completa datos: si una fila no tenía teléfono y su repetida sí, la persona "
        "queda con teléfono. Lo que sigue faltando después de unir:",
        "",
        "| Dato | Personas sin él |",
        "|---|---|",
    ]
    for campo, nombre in CAMPOS.items():
        if faltantes[campo]:
            md.append(f"| {nombre} | {faltantes[campo]} |")
    md += [
        "",
        "## 3. Conciliación con el CRM",
        "",
        "Cada persona se buscó en el CRM con las mismas reglas de los duplicados. Si ya estaba, solo se completa lo "
        "que el CRM tiene vacío (nunca se sobrescribe: el CRM es más reciente que el Excel) y sus notas pasan a "
        "actividades, a nombre del asesor dueño del cliente.",
        "",
        f"- Ya están en el CRM: {len(plan.existentes)}",
        f"- Notas del Excel guardadas en el CRM: {datos.notas_en_crm} (pendientes: {notas_pendientes})",
        f"- Datos vacíos del CRM que el Excel aún puede completar: {datos_pendientes}",
        f"- Diferencias entre el Excel y el CRM (no se cambian): {sum(len(e.diferencias) for e in plan.existentes)}",
        f"- Personas nuevas pendientes de crear: {len(plan.nuevas)}",
        f"- No se pudieron ubicar: {len(casos_del_plan(plan))}",
        "",
    ]
    for caso in casos_del_plan(plan):
        md.append(f"  - **{nombre_de(caso.persona)}** ({donde(caso.persona.filas)}): {caso.razon}")
    md += [
        "",
        "## 4. Verificación: duplicados en el CRM",
        "",
        f"Se compararon los {len(datos.crm)} contactos del CRM entre sí con las mismas reglas: "
        f"**{datos.grupos_duplicados_crm} duplicados**. {estado_crm}",
    ]
    if homonimos:
        md += ["", f"{len(homonimos)} pareja(s) se llaman igual, pero sus datos se contradicen: son homónimos.", ""]
        md += [f"- {p.a.codigo} y {p.b.codigo}, {p.a.nombre}: {p.razon}" for p in homonimos]
    if posibles:
        md += ["", f"{len(posibles)} pareja(s) son posibles duplicados que tiene que revisar una persona:", ""]
        md += [f"- {p.a.codigo} y {p.b.codigo}: {p.razon}" for p in posibles]
    if nota:
        md += [
            "",
            "## 5. Cómo sabemos que está bien",
            "",
            "Como los datos son ficticios, existe un archivo con la respuesta correcta de cada fila "
            "(`data/heredado/respuestas_migracion.csv`). Las reglas no se sacaron de él: solo se usa para calificar.",
            "",
            "| Prueba | Resultado |",
            "|---|---|",
            f"| Datos llenados con un valor equivocado | {nota['equivocados']} |",
            "| Teléfonos recuperados, de las filas donde se podía | {} de {} |".format(*nota["telefonos"]),
            "| Correos recuperados, de las filas donde se podía | {} de {} |".format(*nota["correos"]),
            f"| Fichas que mezclan a dos personas | {nota['mezcladas']} |",
            f"| Personas repartidas en más de una ficha | {nota['repartidas']} |",
            f"| Personas asociadas al contacto correcto del CRM | {nota['bien_asociadas']} de {len(plan.existentes)} |",
        ]
    md += [
        "",
        "## Para revisar",
        "",
        f"La lista completa está en `data/migracion/para_revisar.csv` ({casos} caso(s)).",
        "",
    ]
    return "\n".join(md)


def donde(filas: list[tuple[str, int]]) -> str:
    """[("Medellín (Vale)", 58)] → "Medellín (Vale), fila 58" """
    return "; ".join(f"{hoja}, fila {fila}" for hoja, fila in filas)


def escribir(datos: DatosReporte, nota: dict | None) -> None:
    RUTA_REPORTE.write_text(markdown(datos, nota), encoding="utf-8")
    RUTA_REVISION.parent.mkdir(parents=True, exist_ok=True)
    filas = para_revisar(datos)
    # utf-8-sig: así Excel abre bien las tildes
    with RUTA_REVISION.open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=["tipo", "filas del Excel", "nombre", "motivo"])
        escritor.writeheader()
        escritor.writerows(filas)


def main() -> None:
    from conexion import conectar
    from migrar import leer_crm

    contactos = [normalizar(f) for f in leer_excel()]
    personas, revision_duplicados = deduplicar(contactos)
    with conectar() as con:
        crm, _ = leer_crm(con)
        notas = con.execute(
            "select count(*) from actividades where tipo = 'nota' and descripcion like %s", (PREFIJO_NOTA + "%",)
        ).fetchone()[0]
    grupos_crm, revision_crm = agrupar(crm)
    datos = DatosReporte(
        filas_en_excel=contar_filas_excel(), contactos=contactos, personas=personas,
        revision_duplicados=revision_duplicados, plan=planear(personas, crm), notas_en_crm=notas, crm=crm,
        revision_crm=revision_crm, grupos_duplicados_crm=sum(len(g) > 1 for g in grupos_crm), fecha=date.today(),
    )  # fmt: skip
    se_puede_calificar = RUTA_RESPUESTAS.exists() and RUTA_VERDAD.exists()
    escribir(datos, calificar(datos) if se_puede_calificar else None)
    print(f"Reporte: {RUTA_REPORTE.relative_to(RAIZ)}")
    print(f"Para revisar: {RUTA_REVISION.relative_to(RAIZ)} ({len(para_revisar(datos))} caso(s))")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
