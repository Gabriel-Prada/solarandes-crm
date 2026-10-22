"""Verifica que los datos ficticios sean coherentes, porque todas las fases dependen de ellos."""

import csv
import random
import re
from collections import Counter
from datetime import datetime

import generar_datos as g
import pytest
from openpyxl import load_workbook
from reglas import CIUDADES


@pytest.fixture(scope="module")
def datos():
    return g.generar()


def test_es_reproducible(datos):
    otra_vez = g.generar()
    assert otra_vez.contactos == datos.contactos
    assert otra_vez.oportunidades == datos.oportunidades


def test_contactos_validos(datos):
    assert len(datos.contactos) == g.NUM_CONTACTOS
    for c in datos.contactos:
        assert re.fullmatch(r"\+573\d{9}", c["telefono"])
        assert c["email"].split("@")[1] in g.DOMINIOS_CORREO
        assert c["ciudad"] in CIUDADES
        assert c["asesor_id"] == g.ASESOR_POR_CIUDAD[c["ciudad"]]
        assert c["creado_en"].date() <= g.HOY
    assert len({c["telefono"] for c in datos.contactos}) == g.NUM_CONTACTOS
    assert len({c["email"] for c in datos.contactos}) == g.NUM_CONTACTOS


def test_referencias_existen(datos):
    contactos = {c["id"] for c in datos.contactos}
    oportunidades = {o["id"] for o in datos.oportunidades}
    usuarios = {u["id"] for u in datos.usuarios}
    assert all(o["contacto_id"] in contactos for o in datos.oportunidades)
    assert all(a["oportunidad_id"] in oportunidades for a in datos.actividades)
    assert all(a["usuario_id"] in usuarios for a in datos.actividades)
    assert all(c["oportunidad_id"] in oportunidades for c in datos.cotizaciones)


def test_embudo_realista(datos):
    etapas = Counter(o["etapa"] for o in datos.oportunidades)
    tasa_cierre = etapas["ganado"] / len(datos.oportunidades)
    assert 0.08 <= tasa_cierre <= 0.20
    # Debe haber oportunidades abiertas en todas las etapas para que el tablero se vea real
    for etapa in ("nuevo", "contactado", "visita_agendada", "cotizado", "negociacion"):
        assert etapas[etapa] > 0


def test_oportunidades_coherentes(datos):
    for o in datos.oportunidades:
        assert o["probabilidad"] == g.PROBABILIDAD_ETAPA[o["etapa"]]
        cerrada = o["etapa"] in ("ganado", "perdido")
        assert (o["fecha_cierre"] is not None) == cerrada
        assert bool(o["motivo_perdida"]) == (o["etapa"] == "perdido")
        assert o["creada_en"] <= o["actualizada_en"]


def test_valor_de_oportunidad_es_el_de_su_ultima_cotizacion(datos):
    ultima = {}
    for c in datos.cotizaciones:
        ultima[c["oportunidad_id"]] = c  # están en orden cronológico
    por_id = {o["id"]: o for o in datos.oportunidades}
    for opp_id, c in ultima.items():
        assert por_id[opp_id]["valor_cop"] == c["valor_total"]


def test_cotizaciones(datos):
    numeros = [c["numero"] for c in datos.cotizaciones]
    assert len(set(numeros)) == len(numeros)
    for c in datos.cotizaciones:
        assert re.fullmatch(r"COT-\d{4}-\d{4}", c["numero"])
        assert c["estado"] in ("enviada", "aceptada", "rechazada", "reemplazada")


def test_actividades_en_el_tiempo_correcto(datos):
    ahora = datetime.combine(g.HOY, datetime.max.time())
    for a in datos.actividades:
        if a["completada"]:
            assert a["fecha"] <= ahora, a
        else:
            assert a["fecha"].date() > g.HOY or a["tipo"] in ("visita", "tarea"), a


def test_hay_leads_sin_seguimiento(datos):
    # La fase 3 necesita casos de leads abiertos sin actividad reciente
    ultima = {}
    for a in datos.actividades:
        if a["completada"]:
            ultima[a["oportunidad_id"]] = a["fecha"]
    abiertas = [o for o in datos.oportunidades if o["etapa"] not in ("ganado", "perdido")]
    sin_seguimiento = [o for o in abiertas if (g.HOY - ultima[o["id"]].date()).days > 7]
    assert len(sin_seguimiento) >= 5


@pytest.fixture(scope="module")
def heredado(datos, tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("heredado")
    hojas = g.construir_heredado(random.Random(g.SEMILLA + 1), datos)
    g.escribir_heredado(hojas, carpeta / "h.xlsx", carpeta / "r.csv")
    with (carpeta / "r.csv").open(encoding="utf-8") as f:
        respuestas = list(csv.DictReader(f))
    return load_workbook(carpeta / "h.xlsx"), respuestas


def test_excel_heredado_tiene_las_trampas(heredado):
    libro, respuestas = heredado
    assert libro.sheetnames == [g.HOJA_GENERAL, g.HOJA_MEDELLIN]
    tipos = Counter(r["tipo_fila"] for r in respuestas)
    for trampa in ("vacia", "encabezado_repetido", "total", "titulo"):
        assert tipos[trampa] >= 1
    repetidos = Counter(r["contacto_id"] for r in respuestas if r["tipo_fila"] == "dato")
    assert sum(1 for n in repetidos.values() if n > 1) >= 20


def test_respuestas_coinciden_con_el_excel(heredado, datos):
    libro, respuestas = heredado
    ids = {c["id"] for c in datos.contactos}
    for r in respuestas:
        fila = list(libro[r["hoja"]].iter_rows(min_row=int(r["fila"]), max_row=int(r["fila"]), values_only=True))[0]
        if r["tipo_fila"] == "vacia":
            assert all(v is None for v in fila)
        elif r["tipo_fila"] == "dato":
            assert r["contacto_id"] in ids
            assert any(v is not None for v in fila)
